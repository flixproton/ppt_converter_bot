import io
import logging
import os
import re
import threading
import unicodedata
from http.server import BaseHTTPRequestHandler, HTTPServer
import fitz  # PyMuPDF
from fpdf import FPDF
from pptx import Presentation
from telegram import Update
from telegram.ext import (
    ApplicationBuilder,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)

# -------------------------------------------------------------
# CONFIGURATION & BRANDING
# -------------------------------------------------------------
BRAND_NAME = "Yash PPT Converter Bot"
BOT_TOKEN = os.environ.get("BOT_TOKEN", "YOUR_TELEGRAM_BOT_TOKEN_HERE")
PORT = int(os.environ.get("PORT", 8080))

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)


# -------------------------------------------------------------
# DUMMY SERVER FOR RENDER HEALTH CHECKS
# -------------------------------------------------------------
class HealthCheckHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.send_header("Content-type", "text/plain")
        self.end_headers()
        self.wfile.write(f"{BRAND_NAME} is active!".encode("utf-8"))

    def log_message(self, format, *args):
        return


def run_dummy_server():
    server = HTTPServer(("0.0.0.0", PORT), HealthCheckHandler)
    logger.info(f"Dummy health check server running on port {PORT}")
    server.serve_forever()


# -------------------------------------------------------------
# UNICODE SANITIZER (Fixes the crash with AI/Math/Bullet symbols)
# -------------------------------------------------------------
def clean_text_for_pdf(text: str) -> str:
    """Sanitizes text and replaces unsupported Unicode characters."""
    if not text:
        return ""
    
    replacements = {
        '“': '"', '”': '"', '‘': "'", '’': "'",
        '—': '-', '–': '-', '…': '...', '•': '*',
        '▪': '*', '►': '>', '✔': '/', '✓': '/',
        '→': '->', '←': '<-', '⇒': '=>', '≤': '<=', '≥': '>=',
        '≠': '!=', '±': '+/-', '×': 'x', '÷': '/', '°': ' deg '
    }
    for k, v in replacements.items():
        text = text.replace(k, v)
    
    # Normalize unicode to standard ASCII/Latin characters
    text = unicodedata.normalize('NFKD', text)
    text = text.encode('latin-1', 'ignore').decode('latin-1')
    
    # Strip non-printable control characters
    text = "".join(ch for ch in text if ch.isprintable() or ch in ['\n', '\t'])
    return text.strip()


# -------------------------------------------------------------
# CONTENT EXTRACTION (PPTX & PDF)
# -------------------------------------------------------------
def extract_from_pptx(file_bytes):
    prs = Presentation(io.BytesIO(file_bytes))
    slides_data = []

    for idx, slide in enumerate(prs.slides, start=1):
        title = ""
        if slide.shapes.title and slide.shapes.title.text:
            title = clean_text_for_pdf(slide.shapes.title.text)
        else:
            title = f"Topic / Slide {idx}"

        content = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                for paragraph in shape.text_frame.paragraphs:
                    cleaned = clean_text_for_pdf(paragraph.text)
                    if cleaned and cleaned not in content:
                        content.append(cleaned)

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = clean_text_for_pdf(slide.notes_slide.notes_text_frame.text)

        slides_data.append({
            "num": idx,
            "title": title,
            "content": content,
            "notes": notes,
        })
    return slides_data


def extract_from_pdf(file_bytes):
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    slides_data = []

    for idx, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        raw_lines = [clean_text_for_pdf(line) for line in text.split("\n") if clean_text_for_pdf(line)]

        title = raw_lines[0] if raw_lines else f"Page {idx}"
        content = raw_lines[1:] if len(raw_lines) > 1 else []

        slides_data.append({
            "num": idx,
            "title": title,
            "content": content,
            "notes": "",
        })
    return slides_data


# -------------------------------------------------------------
# BLACK & WHITE PDF GENERATION
# -------------------------------------------------------------
class BWNotesPDF(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 10)
        self.set_text_color(100, 100, 100)
        self.cell(0, 8, f"Study Notes  {BRAND_NAME}", border=0, align="R")
        self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 10, f"Page {self.page_no()} | {BRAND_NAME}", align="C")


def create_bw_text_pdf(slides_data):
    """Generates structured clean B&W notes."""
    pdf = BWNotesPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    for item in slides_data:
        # Avoid empty blocks
        if not item["title"] and not item["content"] and not item["notes"]:
            continue

        pdf.set_draw_color(210, 210, 210)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y())
        pdf.ln(3)

        # Title
        pdf.set_font("Helvetica", "B", 12)
        pdf.set_text_color(0, 0, 0)
        pdf.multi_cell(0, 6, f"[{item['num']}] {item['title']}")
        pdf.ln(2)

        # Content bullets
        if item["content"]:
            pdf.set_font("Helvetica", "", 10)
            pdf.set_text_color(30, 30, 30)
            for bullet in item["content"]:
                pdf.multi_cell(0, 5.5, f"- {bullet}")
            pdf.ln(2)

        # Notes
        if item["notes"]:
            pdf.set_font("Helvetica", "I", 9)
            pdf.set_text_color(70, 70, 70)
            pdf.multi_cell(0, 5, f"Notes: {item['notes']}")
            pdf.ln(2)

        pdf.ln(4)

    return bytes(pdf.output())


def convert_image_pdf_to_bw(file_bytes):
    """Fallback: Converts full slide image decks into clean B&W printable pages."""
    src_doc = fitz.open(stream=file_bytes, filetype="pdf")
    out_doc = fitz.open()

    for page in src_doc:
        pix = page.get_pixmap(colorspace=fitz.csGRAY, dpi=150)
        img_bytes = pix.tobytes("png")
        rect = page.rect
        new_page = out_doc.new_page(width=rect.width, height=rect.height)
        new_page.insert_image(rect, stream=img_bytes)

    return out_doc.tobytes()


# -------------------------------------------------------------
# TELEGRAM BOT HANDLERS
# -------------------------------------------------------------
async def start_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    welcome_text = (
        f"👋 **Welcome to {BRAND_NAME}!**\n\n"
        "📄 Send me any **PPTX** or **PDF presentation**, and I will instantly convert it into a "
        "**clean, ink-saving Black & White study notes PDF**!\n\n"
        "✨ **What I do:**\n"
        "• Strips heavy backgrounds & dark colors\n"
        "• Extracts slide text, bullet points & presenter notes\n"
        "• Formats everything into neat, printable A4 pages\n\n"
        "🚀 **Send your file to start!**"
    )
    await update.message.reply_text(welcome_text, parse_mode="Markdown")


async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    document = update.message.document
    file_name = document.file_name or "presentation"
    file_ext = file_name.split(".")[-1].lower()

    if file_ext not in ["pptx", "pdf"]:
        await update.message.reply_text("⚠️ Please send a valid **.pptx** or **.pdf** file.", parse_mode="Markdown")
        return

    status_msg = await update.message.reply_text("⏳ **Converting into clean B&W notes... Please wait.**", parse_mode="Markdown")

    try:
        # Download file
        tg_file = await context.bot.get_file(document.file_id)
        raw_bytes = await tg_file.download_as_bytearray()
        file_bytes = bytes(raw_bytes)

        output_pdf_bytes = None

        if file_ext == "pptx":
            slides_data = extract_from_pptx(file_bytes)
            output_pdf_bytes = create_bw_text_pdf(slides_data)
        else:
            # Try text extraction first
            slides_data = extract_from_pdf(file_bytes)
            total_text_len = sum(len(s["title"]) + sum(len(c) for c in s["content"]) for s in slides_data)

            # If slide has text, build formatted text PDF. If it's pure images/scanned, convert to grayscale B&W pages
            if total_text_len > 40:
                output_pdf_bytes = create_bw_text_pdf(slides_data)
            else:
                output_pdf_bytes = convert_image_pdf_to_bw(file_bytes)

        # Brand the output filename
        base_name = os.path.splitext(file_name)[0]
        output_filename = f"{base_name}_Yash_PPT_Converter_Bot.pdf"

        # Send result back
        await update.message.reply_document(
            document=io.BytesIO(output_pdf_bytes),
            filename=output_filename,
            caption=f"✅ **Converted successfully by {BRAND_NAME}!**\n📄 Clean Black & White Study Notes.",
            parse_mode="Markdown",
        )
        await status_msg.delete()

    except Exception as e:
        logger.error(f"Error during conversion: {e}", exc_info=True)
        await status_msg.edit_text(f"❌ Conversion failed: `{str(e)}`", parse_mode="Markdown")


# -------------------------------------------------------------
# MAIN APP ENTRY
# -------------------------------------------------------------
def main():
    server_thread = threading.Thread(target=run_dummy_server, daemon=True)
    server_thread.start()

    if not BOT_TOKEN or BOT_TOKEN == "YOUR_TELEGRAM_BOT_TOKEN_HERE":
        print("ERROR: BOT_TOKEN is missing! Set BOT_TOKEN in Render Environment Variables.")
        return

    app = ApplicationBuilder().token(BOT_TOKEN).build()
    app.add_handler(CommandHandler("start", start_command))
    app.add_handler(CommandHandler("help", start_command))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document))

    logger.info(f"{BRAND_NAME} is active and running...")
    app.run_polling()


if __name__ == "__main__":
    main()
